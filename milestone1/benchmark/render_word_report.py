"""Export the verified report content as editable Word text and tables."""
from pathlib import Path
from html.parser import HTMLParser
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import render_report as source

ROOT=Path(__file__).resolve().parents[1]
flow=[]
class Capture:
    def __init__(self,*args,**kwargs): pass
    def build(self,story,**kwargs): flow.extend(story)
source.SimpleDocTemplate=Capture
(ROOT/'tmp/docx').mkdir(parents=True,exist_ok=True)
source.make_pdf(ROOT/'reports/group_benchmark/report_data.json',ROOT/'tmp/docx/source_capture.pdf')
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.5);sec.bottom_margin=Inches(.5)
sec.left_margin=sec.right_margin=Pt(44);sec.footer_distance=Inches(.2)
for name,size in [('Normal',11),('Title',24),('Subtitle',11),('Heading 1',22),('Heading 2',13),('Caption',9.5)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.08
    if name.startswith('Heading'):st.font.bold=True;st.paragraph_format.space_before=Pt(8)
doc.styles['Title'].font.bold=True
foot=sec.footer.paragraphs[0];foot.text='Milestone 1    Final merged dataset    Seed 42';foot.style='Caption'
foot.add_run(' '*8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)

class Inline(HTMLParser):
    def __init__(self,p):super().__init__();self.p=p;self.bold=False;self.italic=False;self.url=None
    def handle_starttag(self,tag,attrs):
        if tag=='b':self.bold=True
        if tag=='i':self.italic=True
        if tag=='br':self.p.add_run().add_break()
        if tag=='link':self.url=dict(attrs).get('href')
    def handle_endtag(self,tag):
        if tag=='b':self.bold=False
        if tag=='i':self.italic=False
        if tag=='link':self.url=None
    def handle_data(self,s):
        if self.url:
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),self.p.part.relate_to(self.url,RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');rp.append(color);r.append(rp)
            t=OxmlElement('w:t');t.text=s;r.append(t);h.append(r);self.p._p.append(h)
        else:
            r=self.p.add_run(s);r.bold=self.bold;r.italic=self.italic

def para(text,style='Normal',p=None):
    p=p if p is not None else doc.add_paragraph(style=style)
    Inline(p).feed(text)
    return p

def table(item):
    values=item._cellvalues;t=doc.add_table(rows=len(values),cols=len(values[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
    pr.append(b)
    for j,w in enumerate(item._colWidths):t.columns[j].width=Pt(w)
    for i,row in enumerate(t.rows):
        if i==0:
            header=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(header)
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
        for j,cell in enumerate(row.cells):
            cell.width=Pt(item._colWidths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'162F46' if i==0 else ('EFF4F7' if i%2 else 'FFFFFF'));cp.append(sh)
            margins=OxmlElement('w:tcMar')
            for edge,v in [('top',70),('bottom',70),('left',100),('right',100)]:
                x=OxmlElement('w:'+edge);x.set(qn('w:w'),str(v));x.set(qn('w:type'),'dxa');margins.append(x)
            cp.append(margins)
            p=para(values[i][j].text,p=cell.paragraphs[0]);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
            if j>0 and len(values[0])!=2 and (re.match(r'^[\d.,% *]+$',p.text) or i==0):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:r.font.size=Pt(9);r.font.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

first=True
pending_break=False
for f in flow:
    if isinstance(f,source.Paragraph):
        text=f.text;name=f.style.name
        if text.startswith('MILESTONE 1'):continue
        if name in ['TitleCustom','SectionCustom']:
            text=text.replace('<br/>',' ').replace(',','').replace('?','')
            style='Title' if first else ('Heading 1' if name=='TitleCustom' else 'Heading 2')
            first=False
        elif name=='SubtitleCustom':text='Milestone 1 report using 20 epochs per model and seed 42';style='Subtitle'
        elif name=='SmallCustom':style='Caption'
        else:style='Normal'
        p=para(text,style)
        if pending_break:
            p.paragraph_format.page_break_before=True;pending_break=False
    elif isinstance(f,source.Table):table(f)
    elif isinstance(f,source.Image):
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(5)
        p.add_run().add_picture(f.filename,width=Pt(f.drawWidth),height=Pt(f.drawHeight))
        pic=p.runs[0]._r.xpath('.//wp:docPr')[0];pic.set('descr',Path(f.filename).stem.replace('_',' '))
    elif isinstance(f,source.PageBreak):pending_break=True
output=ROOT/'output/docx/milestone1_group_comparison.docx';output.parent.mkdir(parents=True,exist_ok=True)
doc.core_properties.title='Milestone 1 Controlled Comparison of Four CNN Models'
doc.core_properties.author='Project group'
for st in doc.styles:
    for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
for p in doc.paragraphs:
    for border in p._p.xpath('.//w:pBdr'):border.getparent().remove(border)
doc.save(output);print(output)
