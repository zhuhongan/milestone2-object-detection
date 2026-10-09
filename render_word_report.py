"""Create the Milestone 2 Word report from verified results and real figures."""
import argparse,json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE=Path(__file__).resolve().parent

def link(p,label,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'0563C1');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def main():
 p=argparse.ArgumentParser();p.add_argument('--data',type=Path,default=HERE/'reports/report_data.json');p.add_argument('--output',type=Path,default=HERE/'reports/milestone2_report.docx');a=p.parse_args();d=json.loads(a.data.read_text());assert d['verified'];plots=a.data.parent/'plots';test=d['metrics']['test'];val=d['metrics']['val'];op=d['operating'];counts=d['audit']['counts'];selected=d['selection']['selected_epoch'];doc=Document();sec=doc.sections[0]
 sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.left_margin=sec.right_margin=Inches(.6);sec.top_margin=sec.bottom_margin=Inches(.5);sec.footer_distance=Inches(.2)
 for name,size in [('Normal',11),('Title',23),('Subtitle',11),('Heading 1',20),('Heading 2',13),('Caption',9.5)]:
  st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.08
  st.font.bold=name in ['Title','Heading 1','Heading 2'];st.font.italic=False
  for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
  if name.startswith('Heading'):st.paragraph_format.space_before=Pt(7)
 footer=sec.footer.paragraphs[0];footer.style='Caption';footer.add_run('Milestone 2    YOLOv8s    Seed 42    ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 def text(s,style='Normal'):return doc.add_paragraph(s,style)
 def heading(s,new=False):
  p=text(s,'Heading 1');p.paragraph_format.page_break_before=new
 def sub(s):text(s,'Heading 2')
 def image(name,width=7.1,caption=False):
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_with_next=caption;p.add_run().add_picture(str(plots/name),width=Inches(width));p.runs[0]._r.xpath('.//wp:docPr')[0].set('descr',name.replace('_',' '))
 def table(headers,rows,widths):
  t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
  for j,w in enumerate(widths):t.columns[j].width=Inches(w)
  for i,rowvalues in enumerate([headers]+rows):
   row=t.rows[0] if i==0 else t.add_row();rp=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');rp.append(no)
   if i==0:rp.append(OxmlElement('w:tblHeader'))
   for j,v in enumerate(rowvalues):
    cell=row.cells[j];cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cp=cell._tc.get_or_add_tcPr();fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'162F46' if i==0 else ('EFF4F7' if i%2 else 'FFFFFF'));cp.append(fill);m=OxmlElement('w:tcMar')
    for e,n in [('top',75),('bottom',75),('left',100),('right',100)]:
     x=OxmlElement('w:'+e);x.set(qn('w:w'),str(n));x.set(qn('w:type'),'dxa');m.append(x)
    cp.append(m);q=cell.paragraphs[0];q.paragraph_format.space_after=Pt(0);q.paragraph_format.line_spacing=1.05;q.add_run(str(v))
    if j>0 and len(headers)>2:q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in q.runs:r.font.size=Pt(9.5);r.font.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
  borders=OxmlElement('w:tblBorders')
  for e in ['top','left','bottom','right','insideH','insideV']:
   x=OxmlElement('w:'+e);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
  t._tbl.tblPr.append(borders);sp=doc.add_paragraph();sp.paragraph_format.space_after=Pt(0);sp.paragraph_format.space_before=Pt(0);sp.paragraph_format.line_spacing=Pt(4);sp.add_run().font.size=Pt(4)

 text('Object detection with YOLOv8','Title');text('Milestone 2 group report using the final merged dataset','Subtitle')
 text(f"We generated 1,460 multi-object images from 365 annotated source photos and trained a YOLOv8s detector for all 73 final Object IDs. The validation-selected checkpoint is epoch {selected}. On 146 test composites, mAP50 is {test['mAP50']*100:.2f}% and mAP50–95 is {test['mAP50_95']*100:.2f}%. These scores describe the generated synthetic dataset.")
 sub('Annotations and dataset generation')
 text('We visually reviewed five source photos per class: three from the fixed training source split, one from validation and one from test. Boxes mark the visible target on the EXIF-oriented primary RGB frame. Pairs of sneakers and two shopping bags receive separate instance boxes. The annotation record contains 365 photos and 371 target boxes; review was performed by an assistant and has not been independently verified by a human.')
 table(['Split','Composites','Target instances','Source photos'],[[s.title(),f"{counts[s]['images']:,}",f"{counts[s]['target_instances']:,}",counts[s]['unique_source_photos']] for s in ['train','val','test']],[2.2,1.5,1.8,1.6])
 text('The image split is 80/10/10 with seed 42. Sources come from the finalized, exact-deduplicated Milestone 1 manifests and are separated before synthesis. Paths, byte hashes and RGB-pixel hashes are disjoint across splits. Each split contains all 73 classes. Ten unsuitable annotation candidates were replaced within their original source split; filename-tagged background candidates were ineligible for this selected annotation pool. The shared dataset was preserved.')
 text('Each 640 × 640 composite combines 2–4 randomly selected distinct classes, with 2–5 labeled instances when a source contains a pair. We crop each annotated target with 5–15% context, keep its aspect ratio, resize it and place it at a random non-overlapping location. Crowded layouts can reduce patch size. Training patches receive mild brightness/contrast changes and horizontal flips. Crop, scale and placement transforms update every box; labels use normalized YOLO center, width and height coordinates.')
 image('generation_examples.png',6.2,True);text('Two generated training composites. Red boxes are annotations, not model detections. Context inside each rectangular photo crop is retained.','Caption')

 heading('Training and measured performance',True)
 text('We fine-tuned COCO-pretrained YOLOv8s with Ultralytics 8.3.221 through Colab CLI on an NVIDIA A100-SXM4-40GB. Every planned epoch completed. RGB input tensors are divided by 255 to scale pixels to 0–1. The best checkpoint was chosen by validation mAP50–95 before the benchmark evaluated test images.')
 table(['Setting','Protocol'],[['Budget and input','50 epochs; 640 pixels; batch 32; seed 42; no early stopping'],['Optimization','AdamW; LR 0.001; cosine final factor 0.01; weight decay 0.0005; 3 warm-up epochs'],['Augmentation','Training-only mosaic; closed for final 10 epochs; rotation ±10°; translation 0.1; scale 0.3; mild HSV changes; horizontal flip 0.5'],['Evaluation','Fixed generated images; IoU thresholds 0.50–0.95 for AP; NMS IoU 0.7']],[1.5,5.6])
 table(['Split','Precision','Recall','mAP50','mAP50–95'],[[name.title(),f"{r['precision']*100:.2f}%",f"{r['recall']*100:.2f}%",f"{r['mAP50']*100:.2f}%",f"{r['mAP50_95']*100:.2f}%"] for name,r in [('val',val),('test',test)]],[1.5,1.4,1.4,1.4,1.4])
 text('AP measures the precision–recall curve for each class. mAP50 requires matching boxes at IoU ≥ 0.50; mAP50–95 averages ten IoU thresholds and rewards tighter localization. Table precision/recall are library class means at the best mean-F1 point on each split’s curve. The fixed-threshold results below use only validation to choose confidence.','Caption')
 image('learning_curves.png',6.8)
 text(f"At a confidence threshold of {op['confidence']:.2f}, chosen only on validation micro-F1, the test set has {op['tp']} true positives, {op['fp']} false positives and {op['fn']} missed targets. Class-aware matching requires IoU ≥ 0.50. Micro precision is {op['precision']*100:.2f}%, recall {op['recall']*100:.2f}% and F1 {op['f1']:.4f}. All annotated objects are correctly detected with no false positives in {op['images_with_all_objects_correct']}/{op['image_count']} images ({op['all_objects_accuracy']*100:.2f}%).")

 heading('Sample detection results',True)
 text('Four of six test paths sampled with seed 42 are shown in their original sample order. Paths were selected independently of predictions. Teal boxes and labels are model outputs at the fixed validation-selected confidence threshold. TP, FP and FN use the same class-aware IoU matching as the reported operating metrics.','Caption')
 image('test_examples.png',6.7,True)
 text('Labels show the final Object ID and model confidence. Box coordinates in the saved CSV/JSON refer to original-image pixels, with the origin at the upper-left. Every test image contributes to the aggregate metrics; this gallery is illustrative.','Caption')

 heading('Error analysis and interpretation',True)
 text(f"Validation mAP50–95 exceeds test mAP50–95 by {(val['mAP50_95']-test['mAP50_95'])*100:.2f} percentage points. The test images use different source photos, so new viewpoints, partial views and retained background context can change detection quality. This gap is evidence of limited generalization within the synthesis setup; it does not identify a single cause.")
 image('per_class_ap.png',6.6)
 text('Per-class AP uses a small number of repeated views derived from one test source photo per class. Individual class scores can therefore be unstable. The full 73-class table is archived.','Caption')
 text('The case below was selected by most missed targets, then false positives and filename. Red boxes show ground truth; teal boxes show predictions. This example explains a failure and does not estimate its frequency.','Caption')
 image('failure_example_one.png',6.7)

 heading('Final assessment and reproducibility',True)
 text(f"The trained YOLOv8s model detects Object IDs and locations on the held-out synthetic composites, with test mAP50–95 {test['mAP50_95']*100:.2f}%. The strict all-objects image score is {op['all_objects_accuracy']*100:.2f}%, so the model does not reliably find every target in every image. We retain the validation-selected checkpoint and confidence threshold for inference.")
 sub('Example pixel locations')
 name=d['sample_files'][0];text(f"The following detections are from {name}, the first independently sampled example. Each coordinate tuple is (x1, y1, x2, y2) in its 640 × 640 image.",'Caption')
 coord=d['sample_coordinates'];table(['Predicted ID','Official object name','Confidence','Box in pixels'],[[r['object_id'],r['object_name'],f"{float(r['confidence']):.2f}",', '.join(str(round(float(r[k]))) for k in ['x1','y1','x2','y2'])] for r in coord],[1.15,2.6,1.0,2.35])
 text('OBJ018 is a correct detection. The OBJ055 prediction is incorrect: the annotated target is OBJ020 Airpod. This yields one false positive and one missed target in this example.','Caption')
 sub('Limits and next steps')
 text('Only three distinct training source photos and one validation/test source photo were annotated per class. Repeated synthetic crops increase the number of images but not the number of independent photographs. Exact deduplication does not establish independence between nearby captures or the same physical object. Rectangular crop borders, retained backgrounds, truncated objects and approximate annotations can affect learning. The initial annotation pool also omits background-tagged and unsuitable candidates, so it represents a selected subset of the finalized dataset.')
 text('We have not tested on independently photographed real multi-object scenes. A stronger next experiment would use more annotated sources per class, an independently reviewed annotation set, realistic occlusion and backgrounds, and a separate real-scene test collection. Results from the current synthetic benchmark should not be presented as real-scene accuracy.')
 sub('Code and experiment record')
 q=text('GitHub repository: ');link(q,d['repository'],d['repository'])
 text('The private repository contains generation, validation, training, inference and reporting code, final IDs, reviewed source annotations and this report. The dataset remains outside Git. Training histories, package versions, checkpoint hashes, generated manifests, per-class metrics and all test detections are saved with the experiment. Local geometry, matching and seeded-generation tests passed, and all 1,460 generated image/label pairs passed validation. Colab artifacts are downloaded in checksum-verified parts before the GPU is released.')
 text('Inference: run predict.py with best.pt, your image and operating_threshold.json. Its annotated image and detections.json provide final IDs, official names, confidences and pixel locations.','Caption')
 sub('References')
 for label,url in [('Final merged dataset','https://drive.google.com/drive/folders/1GKLyP17tm5JC0XENv25GeoGgYynm97LS'),('Final Object ID sheet','https://docs.google.com/spreadsheets/d/1OCBWlTK8JSzrgUxON3thJWAKRKU2Pi-V9Ssn2BbDxds/edit#gid=0'),('Ultralytics YOLOv8 model documentation','https://docs.ultralytics.com/models/yolov8/'),('Ultralytics YOLO detection dataset format','https://docs.ultralytics.com/datasets/detect/')]:link(text('', 'Caption'),label,url)
 a.output.parent.mkdir(parents=True,exist_ok=True);doc.core_properties.title='Milestone 2 Object Detection with YOLOv8';doc.core_properties.author='Project group';doc.save(a.output);print(a.output)
if __name__=='__main__':main()
